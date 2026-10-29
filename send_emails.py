#!/usr/bin/env python3
"""
Job-application mailer + Excel tracker.

  python send_emails.py                       -> process every pending row in "To Send"
  python send_emails.py --dry-run             -> show what would happen, send nothing
  python send_emails.py --add "Acme" hr@acme.com "Mrs Dupont"   -> add a row to "To Send"

Everything you edit lives OUTSIDE this file:
  config.json   (your address, CV path, mode, delay)
  template.txt  (subject + message, use {company} and {contact})
  tracker.xlsx  (sheet "To Send" = your queue, sheet "Sent Log" = history)
"""
import argparse
import imaplib
import json
import smtplib
import ssl
import sys
import time
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

BASE = Path(__file__).resolve().parent
TO_SEND = "To Send"
LOG = "Sent Log"
QUEUE_HEADERS = ["Company", "Email", "Contact name (optional)", "Status"]
LOG_HEADERS = ["Date", "Month", "Company", "Email", "Subject", "Message", "Mode", "Opened? (from extension)"]


# ---------- config / template ----------
def load_config():
    with open(BASE / "config.json", encoding="utf-8") as f:
        return json.load(f)


def load_template(cfg):
    text = (BASE / cfg["template_file"]).read_text(encoding="utf-8")
    first, _, rest = text.partition("\n")
    if not first.lower().startswith("subject:"):
        sys.exit("template.txt must start with a line like:  Subject: ...")
    return first.split(":", 1)[1].strip(), rest.strip()


class SafeDict(dict):
    def __missing__(self, key):  # unknown {placeholder} stays as is
        return "{" + key + "}"


# ---------- Excel ----------
def style_header(ws, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = "A2"


def ensure_tracker(path):
    if path.exists():
        return
    wb = Workbook()
    q = wb.active
    q.title = TO_SEND
    style_header(q, QUEUE_HEADERS, [28, 34, 26, 22])
    # Example row (ignored by the script because of example.com) - delete it
    q.append(["Example Company", "recruiter@example.com", "Mrs Example", ""])
    lg = wb.create_sheet(LOG)
    style_header(lg, LOG_HEADERS, [14, 16, 28, 34, 40, 80, 10, 22])
    wb.save(path)
    print(f"Created {path.name}")


def open_tracker(path):
    ensure_tracker(path)
    try:
        return load_workbook(path)
    except PermissionError:
        sys.exit("Close tracker.xlsx in Excel first, then run the script again.")


def save_tracker(wb, path):
    try:
        wb.save(path)
    except PermissionError:
        sys.exit("Close tracker.xlsx in Excel first (it is locked), then run again.")


def log_row(wb, company, to, subject, body, mode):
    now = datetime.now()
    ws = wb[LOG]
    ws.append([now, now.strftime("%B %Y"), company, to, subject, body, mode, ""])
    r = ws.max_row
    ws.cell(row=r, column=1).number_format = "yyyy-mm-dd hh:mm"
    ws.cell(row=r, column=6).alignment = Alignment(wrap_text=True, vertical="top")


# ---------- email ----------
def build_message(cfg, to, subject, body):
    msg = EmailMessage()
    msg["From"] = f'{cfg["sender_name"]} <{cfg["email"]}>'
    msg["To"] = to
    msg["Subject"] = subject
    msg.set_content(body)
    cv = Path(cfg["cv_path"])
    if not cv.is_absolute():
        cv = BASE / cv
    if not cv.exists():
        sys.exit(f"CV not found: {cv}")
    msg.add_attachment(cv.read_bytes(), maintype="application", subtype="pdf" if cv.suffix.lower() == ".pdf" else "octet-stream", filename=cv.name)
    return msg


def send_smtp(cfg, msg):
    with smtplib.SMTP_SSL("smtp.gmail.com", 465, context=ssl.create_default_context()) as s:
        s.login(cfg["email"], cfg["app_password"])
        s.send_message(msg)


def drafts_folder(imap, cfg):
    if cfg.get("drafts_folder"):
        return cfg["drafts_folder"]
    _, boxes = imap.list()
    for b in boxes:
        line = b.decode(errors="ignore")
        if "\\Drafts" in line:
            return line.split(' "/" ')[-1].strip().strip('"')
    return "[Gmail]/Drafts"


def save_draft(cfg, msg):
    imap = imaplib.IMAP4_SSL("imap.gmail.com")
    try:
        imap.login(cfg["email"], cfg["app_password"])
        folder = drafts_folder(imap, cfg)
        imap.append(folder, "\\Draft", imaplib.Time2Internaldate(time.time()), msg.as_bytes())
    finally:
        try:
            imap.logout()
        except Exception:
            pass


# ---------- main ----------
def add_row(args):
    cfg = load_config()
    path = BASE / cfg["tracker_file"]
    wb = open_tracker(path)
    wb[TO_SEND].append([args[0], args[1], args[2] if len(args) > 2 else "", ""])
    save_tracker(wb, path)
    print(f"Added {args[0]} <{args[1]}> to '{TO_SEND}'")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--add", nargs="+", metavar=("COMPANY", "EMAIL"))
    a = ap.parse_args()

    if a.add:
        if len(a.add) < 2:
            sys.exit('Usage: --add "Company" email@company.com ["Contact name"]')
        return add_row(a.add)

    cfg = load_config()
    mode = cfg.get("mode", "draft")  # "draft" or "send"
    subject_t, body_t = load_template(cfg)
    path = BASE / cfg["tracker_file"]
    wb = open_tracker(path)
    ws = wb[TO_SEND]

    done = 0
    for r in range(2, ws.max_row + 1):
        company, to, contact, status = (ws.cell(row=r, column=c).value for c in range(1, 5))
        if not to or status or str(to).lower().endswith("@example.com"):
            continue
        to = str(to).strip()
        fields = SafeDict(company=company or "", contact=contact or cfg.get("default_contact", ""))
        subject = subject_t.format_map(fields)
        body = body_t.format_map(fields)

        if a.dry_run:
            print(f"[dry-run] {mode} -> {company} <{to}> | {subject}")
            continue

        msg = build_message(cfg, to, subject, body)
        try:
            if mode == "send":
                send_smtp(cfg, msg)
            else:
                save_draft(cfg, msg)
        except Exception as e:  # keep going, mark the failure
            ws.cell(row=r, column=4, value=f"ERROR: {e}")
            save_tracker(wb, path)
            print(f"FAILED {company} <{to}>: {e}")
            continue

        ws.cell(row=r, column=4, value="Sent" if mode == "send" else "Draft created - press Send in Gmail")
        log_row(wb, company, to, subject, body, mode)
        save_tracker(wb, path)  # save after each mail: no duplicates if it crashes
        done += 1
        print(f"OK ({mode}) {company} <{to}>")
        time.sleep(cfg.get("delay_seconds", 20))

    print(f"Done: {done} processed.")


if __name__ == "__main__":
    main()
